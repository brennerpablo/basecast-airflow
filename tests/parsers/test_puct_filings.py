import io
import json
import zipfile
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from basecast_pipelines.parsers.puct.interchange import (
    DATASETS,
    _page_ranges,
    classify_party,
    parse_documents,
    parse_pages,
    parse_tsp_requests,
    rows_from_grid,
    scrub_names,
)
from tests.parsers.pdf_fixture import minimal_pdf, table_page, text_page

SOURCE = "puct_filings"
DT = date(2026, 9, 26)

# Attachment A, slide 5 (by TSP), trimmed to four TSPs; cells as pdfplumber reads the PDF export.
TSP_GRID = [
    ["TSP", "2026", "2027", "2028", "2029"],
    ["AEP", "2,410", "8,682", "17,536", "22,944"],
    ["Rayburn", "-", "224", "845", "1,446"],
    ["STEC\n(South Texas)", "-", "120", "274", "545"],
    ["Aggregate of\nTSPs with < 3\nsites", "297", "307", "1,617", "3,491"],
    ["Total", "2,707", "9,333", "20,272", "28,426"],
]
TYPE_GRID = [
    ["Load Type", "2026", "2027", "2028"],
    ["Data Center (non-\ncrypto)", "7,401", "39,719", "99,350"],
    ["Cryptocurrency\nMining", "2,475", "3,518", "5,611"],
    ["Total", "9,876", "43,237", "104,961"],
]


@pytest.mark.parametrize("party, org, individual", [
    ("ONCOR ELECTRIC DELIVERY COMPANY LLC", "ONCOR ELECTRIC DELIVERY COMPANY LLC", False),
    ("Cholla, Inc.", "Cholla, Inc.", False),
    ("SHELL ENERGY NORTH AMERICA, (US), L.P.", "SHELL ENERGY NORTH AMERICA, (US), L.P.", False),
    ("GOOGLE LLC, LANCIUM LLC", "GOOGLE LLC, LANCIUM LLC", False),
    ("TIEC", "TIEC", False),
    ("AEP COMPANIES", "AEP COMPANIES", False),
    ("Great Bay Royalties", "Great Bay Royalties", False),
    ("JANE Q PUBLIC", None, True),
    ("John Doe", None, True),
    ("State Senator John Doe", None, True),
    ("John Doe, Example Labs", "Example Labs", True),
    ("John Doe / Example AI", "Example AI", True),
])
def test_filing_party_keeps_organizations_only(party, org, individual):
    assert classify_party(party)[:2] == (org, individual)


def test_individual_names_are_scrubbed_from_descriptions():
    _, _, persons = classify_party("John Doe, Example Labs")
    text = "Comments of John Doe on SB 6; Doe also asks for a hearing. State of Texas."
    assert scrub_names(text, persons) == "Comments of [individual] on SB 6; [individual] also asks for a hearing. State of Texas."


def test_rows_from_grid_reads_years_dashes_and_totals():
    rows = rows_from_grid(TSP_GRID)
    by = {(name, year): mw for name, year, mw, _ in rows}
    assert by[("Rayburn", 2026)] == 0 and by[("STEC (South Texas)", 2027)] == 120
    assert by[("Aggregate of TSPs with < 3 sites", 2029)] == 3491
    totals = {year: mw for name, year, mw, is_total in rows if is_total}
    sums = {}
    for name, year, mw, is_total in rows:
        if not is_total:
            sums[year] = sums.get(year, 0) + mw
    assert totals == sums == {2026: 2707, 2027: 9333, 2028: 20272, 2029: 28426}
    assert {n for n, *_ in rows_from_grid(TYPE_GRID)} == {"Data Center (non-crypto)", "Cryptocurrency Mining", "Total"}
    with pytest.raises(ValueError):
        rows_from_grid([["TSP", "2026", "2027", "2028"], ["AEP", "2,410", "n/a?", "1"]])


def _attachment_pptx() -> bytes:
    prs = Presentation()
    for title, grid in (("TSP Large Load RFI Submissions by Load Type", TYPE_GRID),
                        ("TSP Large Load RFI Submissions by TSP", TSP_GRID)):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        table = slide.shapes.add_table(len(grid), len(grid[0]), Inches(0.5), Inches(1.5), Inches(9), Inches(4)).table
        for r, row in enumerate(grid):
            for c, cell in enumerate(row):
                table.cell(r, c).text = f"  {cell} " if r and c else cell  # the real cells are space-padded
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _attachment_pdf() -> bytes:
    flat = [[c.replace("\n", " ") for c in row] for row in TYPE_GRID]
    flat_tsp = [[c.replace("\n", " ") for c in row] for row in TSP_GRID]
    return minimal_pdf([
        table_page("TSP Large Load RFI Submissions by Load Type", flat),
        table_page("TSP Large Load RFI Submissions by TSP", flat_tsp, col_widths=[200, 70, 70, 70, 70]),
    ])


@pytest.fixture
def index(raw_file):
    return raw_file(SOURCE, "puct_filings/58777_filings_trimmed.xlsx", dt=DT, name="58777_filings.xlsx",
                    meta={"control_number": 58777, "kind": "filings_index"})


def _zip(tmp_path: Path, name: str, members: dict[str, bytes]) -> Path:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        for member, data in members.items():
            z.writestr(member, data)
    path = tmp_path / name
    path.write_bytes(buffer.getvalue())
    return path


def test_tsp_requests_from_pptx_and_matching_pdf(raw_file, index, tmp_path: Path):
    path = _zip(tmp_path, "58777_38_1622646.ZIP", {
        "2026-04-15 (58777) Preliminary Long-Term Load Forecast.docx": b"not parsed",
        "Attachment A.pptx": _attachment_pptx(),
        "Attachment A.pdf": _attachment_pdf(),
    })
    f = raw_file(SOURCE, path, dt=DT, meta={"control_number": 58777, "item": 38, "kind": "document"})
    df = parse_tsp_requests(f)
    assert set(df["extraction_method"]) == {"pptx_table", "pdf_table"}
    assert (df["docket"][0], df["item"][0], df["document_id"][0], df["filed_date"][0]) == ("58777", 38, "1622646", date(2026, 4, 15))
    pptx = df.filter(df["extraction_method"] == "pptx_table")
    pdf = df.filter(df["extraction_method"] == "pdf_table")
    key = ["breakdown", "name", "year", "mw", "is_total"]
    assert pptx.select(key).sort(key).equals(pdf.select(key).sort(key))
    totals = pptx.filter(pptx["is_total"] & (pptx["breakdown"] == "tsp")).sort("year")["mw"].to_list()
    assert totals == [2707, 9333, 20272, 28426]
    assert set(df["verified"]) == {False}


def test_other_zips_have_no_tsp_table(raw_file, index, tmp_path: Path):
    path = _zip(tmp_path, "58481_100_1575739.ZIP", {"Comments.pdf": minimal_pdf([text_page(["Comments"])])})
    f = raw_file(SOURCE, path, dt=DT, meta={"control_number": 58481, "item": 100, "kind": "document"})
    assert parse_tsp_requests(f) is None


def _rendition(lines: list[str]) -> Path:
    return minimal_pdf([
        text_page(["Filing Receipt", "Filed Date - 2026-04-15 03:47:02 PM", "Control Number - 58777", "Item Number - 38"]),
        text_page(lines),
        {"lines": []},  # a scanned page: no text layer
    ])


def test_documents_and_pages_of_an_organization_filing(raw_file, index, tmp_path: Path):
    path = tmp_path / "58777_38_1622647.PDF"
    path.write_bytes(_rendition(["ERCOT files the preliminary LTLF.", "Questions: someone@ercot.com, T: 512-248-3000"]))
    f = raw_file(SOURCE, path, dt=DT, meta={"control_number": 58777, "item": 38, "kind": "document"})
    doc = parse_documents(f).row(0, named=True)
    assert doc["filing_party"] == "ERCOT" and doc["filed_by_individual"] is False
    assert doc["filed_date"] == date(2026, 4, 15) and doc["description"].startswith("PRELIMINARY LONG-TERM")
    assert (doc["pages"], doc["pages_extracted"], doc["pages_without_text"], doc["pages_without_text_list"]) == (3, 3, 1, "3")
    assert doc["has_text_layer"] and doc["text_stored"] and doc["error"] is None and doc["file_type"] == "pdf"
    pages = parse_pages(f)
    assert pages["page"].to_list() == [1, 2]
    assert "[email]" in pages["text"][1] and "[phone]" in pages["text"][1] and "248-3000" not in pages["text"][1]
    assert (pages["docket"][0], pages["item"][0], pages["document_date"][0]) == ("58777", 38, date(2026, 4, 15))


def test_filings_by_individuals_keep_no_name_and_no_text(raw_file, index, tmp_path: Path):
    path = tmp_path / "58777_40_1638426.PDF"
    path.write_bytes(_rendition(["Letter from Jane Q Public, 1 Main St."]))
    f = raw_file(SOURCE, path, dt=DT, meta={"control_number": 58777, "item": 40, "kind": "document"})
    doc = parse_documents(f).row(0, named=True)
    assert doc["filing_party"] is None and doc["filed_by_individual"] is True and doc["text_stored"] is False
    assert "Jane" not in doc["description"] and "Public" not in doc["description"]
    assert parse_pages(f) is None


def test_zip_documents_count_members_and_read_pptx(raw_file, index, tmp_path: Path):
    path = _zip(tmp_path, "58777_38_1622646.ZIP", {
        "Attachment A.pptx": _attachment_pptx(),
        "Attachment A.pdf": _attachment_pdf(),
        "cover.docx": b"x",
    })
    f = raw_file(SOURCE, path, dt=DT, meta={"control_number": 58777, "item": 38, "kind": "document"})
    doc = parse_documents(f).row(0, named=True)
    assert doc["zip_members"] == 3 and json.loads(doc["zip_member_types"]) == {"docx": 1, "pdf": 1, "pptx": 1}
    assert doc["pages"] == 4 and doc["pages_extracted"] == 2 and doc["text_stored"]
    pages = parse_pages(f)
    assert set(pages["member"]) == {"Attachment A.pptx"}
    assert pages["page_title"].to_list() == ["TSP Large Load RFI Submissions by Load Type", "TSP Large Load RFI Submissions by TSP"]


def test_page_ranges():
    assert _page_ranges([10, 41, 62, 63, 64, 66]) == "10, 41, 62-64, 66"
    assert _page_ranges([]) is None


def test_datasets_are_declared():
    assert {d.name: d.mode for d in DATASETS} == {
        "puct_tsp_large_load_requests": "by_file", "puct_filing_documents": "by_file", "document_pages": "by_file"}
