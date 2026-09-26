import io
import zipfile
from datetime import date
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from basecast_pipelines.parsers._documents import scrub
from basecast_pipelines.parsers.ercot.large_load import (
    DATASETS,
    decks,
    definition_version,
    headlines,
    parse_headlines,
    parse_pages,
    parse_status,
)
from tests.parsers.conftest import FIXTURES
from tests.parsers.pdf_fixture import minimal_pdf, text_page

SOURCE = "ercot_large_load_decks"
PAGES = yaml.safe_load((FIXTURES / SOURCE / "page_texts.yaml").read_text())


def rows_for(name: str) -> list[dict]:
    page = PAGES[name]
    return headlines(page["text"], report_date=page["report_date"])


def one(rows: list[dict], metric: str, **match) -> dict:
    found = [r for r in rows if r["metric"] == metric and all(r[k] == v for k, v in match.items())]
    assert len(found) == 1, (metric, match, rows)
    return found[0]


def test_monthly_peaks_and_approved_load():
    rows = rows_for("march_2026_updated_p6")
    assert one(rows, "approved_to_energize_mw")["value"] == 9042
    peak = one(rows, "observed_peak_non_simultaneous_mw")
    assert (peak["value"], peak["peak_basis"], peak["as_of_date"], peak["as_of_suspect"]) == (
        4004, "monthly", date(2026, 3, 1), False)
    assert "4,004 MW in March 2026" in peak["text_snippet"]
    simultaneous = one(rows_for("march_2026_updated_p7"), "observed_peak_simultaneous_mw")
    assert (simultaneous["value"], simultaneous["peak_basis"]) == (3522, "monthly")


def test_wrong_year_in_the_deck_is_flagged():
    """The 2026-03-12 PDF printed "3883 MW in March 2025" (ERCOT's updated deck says 4,004 MW in March 2026)."""
    peak = one(rows_for("march_2026_original_p6"), "observed_peak_non_simultaneous_mw")
    assert (peak["value"], peak["as_of_date"], peak["as_of_suspect"]) == (3883, date(2025, 3, 1), True)


def test_all_time_peak_and_definition_based_basis():
    rows = rows_for("jan_2025_p7")
    monthly = one(rows, "observed_peak_non_simultaneous_mw", peak_basis="monthly")
    all_time = one(rows, "observed_peak_non_simultaneous_mw", peak_basis="all_time")
    assert (monthly["value"], monthly["as_of_date"]) == (3211, date(2025, 1, 1))
    assert (all_time["value"], all_time["as_of_date"], all_time["as_of_suspect"]) == (3697, date(2024, 11, 1), False)
    # 2023 decks give no month; the page defines the number as "regardless of when that maximum occurred".
    old = one(rows_for("may_2023_p4"), "observed_peak_non_simultaneous_mw")
    assert (old["value"], old["peak_basis"], old["as_of_date"]) == (2072, "all_time", None)


def test_breakdowns_by_load_zone_and_project_type():
    rows = rows_for("oct_2024_p5")
    assert one(rows, "approved_to_energize_mw", dimension=None)["value"] == 5697
    by_zone = {r["category"]: r["value"] for r in rows if r["dimension"] == "load_zone"}
    by_type = {r["category"]: r["value"] for r in rows if r["dimension"] == "project_type"}
    assert by_zone == {"LZ_WEST": 3055, "Other": 2642}
    assert by_type == {"Standalone": 4622, "Co-Located": 1075}


def test_tracked_queue_in_gw_with_dates_and_share():
    rows = rows_for("board_sep_2025_p2")
    current = one(rows, "total_mw_tracked", as_of_date=date(2025, 9, 2))
    assert (current["value"], current["unit"], current["qualifier"]) == (189_000, "MW", "approx")
    assert one(rows, "total_mw_tracked", as_of_date=date(2024, 9, 1))["value"] == 56_000
    assert one(rows, "data_center_share_of_tracked_pct")["value"] == 69


def test_ercot_monthly_paragraphs():
    rows = rows_for("monthly_nov_2025_p8")
    assert one(rows, "total_mw_tracked", as_of_date=date(2025, 11, 18))["value"] == 226_000
    assert one(rows, "total_mw_tracked", as_of_date=date(2024, 12, 1))["value"] == 63_000
    assert one(rows, "no_studies_submitted_mw")["value"] == 128_000
    assert one(rows, "approved_to_energize_mw")["value"] == 7_500
    assert one(rows, "observed_energized_mw")["value"] == 5_300
    assert one(rows, "approved_not_operational_mw")["value"] == 2_200
    april = rows_for("monthly_apr_2026_p4")
    assert {r["metric"]: r["value"] for r in april} == {
        "total_mw_tracked": 445_800, "no_studies_submitted_mw": 321_000, "under_ercot_review_mw": 93_700,
        "section_9_5_requirements_met_mw": 22_000, "observed_energized_mw": 5_900, "approved_not_operational_mw": 3_200,
    }


def test_batch_zero_classes():
    rows = rows_for("batch_zero_sep_2026_p1")
    assert {r["metric"]: r["value"] for r in rows} == {
        "batch_zero_base_load_mw": 66_400, "batch_zero_base_load_projects": 204,
        "batch_zero_studied_load_mw": 127_900, "batch_zero_studied_load_projects": 158,
    }


@pytest.mark.parametrize("name", ["jan_2023_p9_garbled", "dec_2023_p4"])
def test_broken_or_unknown_sentences_yield_nothing(name):
    assert rows_for(name) == []


def test_scrub_removes_emails_and_phones_but_not_numbers():
    text = "Contact someone@ercot.com or (512) 248-3000 / 512.225.7000; queue 2,475 3,518 and 250 500 1000 MW"
    assert scrub(text) == "Contact [email] or [phone] / [phone]; queue 2,475 3,518 and 250 500 1000 MW"


def test_definition_version():
    assert definition_version("Section 9.4/9.5 Requirements Met - Projects ...") == "section_9_4_9_5"
    assert definition_version("Observed Energized - Projects ...") == "observed_energized"
    assert definition_version("Planning Studies Approved - ...") == "planning_studies_approved"
    assert definition_version("Met Planning - ...") == "met_planning"


# --- files ------------------------------------------------------------------------------------------


def _deck_pdf() -> bytes:
    return minimal_pdf([
        text_page(["Large Load Interconnection Status Update", "Large Load Integration Team", "March 26, 2026"]),
        text_page([
            "Loads Approved to Energize - Observations 1",
            "Of the 9,042 MW that have received Approval to Energize, ERCOT has observed a",
            "non-simultaneous monthly peak consumption of 4,004 MW in March 2026.",
            "Questions? email: someone@ercot.com",
        ]),
    ])


def test_pdf_deck_headlines_and_pages(raw_file, tmp_path: Path):
    path = tmp_path / "March-TAC-Report.pdf"
    path.write_bytes(_deck_pdf())
    f = raw_file(SOURCE, path, dt=date(2026, 3, 13), meta={"type": "status_deck", "link_text": "March TAC Report",
                                                            "posted_date": "2026-03-12"})
    df = parse_headlines(f)
    assert df.height == 2
    peak = df.filter(df["metric"] == "observed_peak_non_simultaneous_mw").row(0, named=True)
    assert peak["value"] == 4004 and peak["page"] == 2 and peak["extraction_method"] == "pdf_text"
    assert peak["deck_date"] == date(2026, 3, 26) and peak["report_date"] == date(2026, 3, 13)
    assert peak["posted_date"] == date(2026, 3, 12) and peak["verified"] is False
    pages = parse_pages(f)
    assert pages["page"].to_list() == [1, 2]
    assert "[email]" in pages["text"][1] and "@" not in pages["text"][1]
    assert pages["page_title"][1] == "Loads Approved to Energize - Observations 1"
    assert set(pages["truncated"]) == {False} and pages["source_id"][0] == SOURCE


def _chart_deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "ERCOT Approvals - Past 12 Months"
    data = CategoryChartData(number_format="#,##0")
    data.categories = [date(2024, 9, 6), date(2024, 10, 30)]
    data.add_series("Approved to Energize", (5496, 5697))
    data.add_series("Planning Studies Approved", (8754, 12401))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(1), Inches(1.5), Inches(8), Inches(4), data).chart
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Load Amount (MW)"
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.text = "Snapshot Date"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Batch Zero Eligibility"
    data = CategoryChartData()
    data.categories = ["LZ_WEST", "Other"]
    data.add_series("Observed Non-Simultaneous Peak", (1.852, 1.662))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(1), Inches(1.5), Inches(8), Inches(4), data).chart
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Large Load Requests (GW)"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_native_pptx_charts_inside_a_tac_zip(raw_file, tmp_path: Path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("LLI Queue Status Update - 2024-10-30.pptx", _chart_deck())
        z.writestr("Brownsville_Area_Improvement_EIR_TAC_October_2024.pptx", _chart_deck())  # not a large-load deck
    path = tmp_path / "15-ercot-reports.zip"
    path.write_bytes(buffer.getvalue())
    f = raw_file(SOURCE, path, dt=date(2024, 10, 30),
                 meta={"type": "status_deck", "zip_members": ["LLI Queue Status Update - 2024-10-30.pptx"]})
    assert [d.member for d in decks(f)] == ["LLI Queue Status Update - 2024-10-30.pptx"]
    df = parse_status(f)
    assert df.height == 6
    approvals = df.filter(df["status_bucket"] == "approved_to_energize").sort("snapshot_date")
    assert approvals["snapshot_date"].to_list() == [date(2024, 9, 6), date(2024, 10, 30)]
    assert approvals["mw"].to_list() == [5496, 5697] and set(approvals["dimension"]) == {"snapshot_date"}
    zones = df.filter(df["page"] == 2).sort("load_zone")
    assert zones["load_zone"].to_list() == ["LZ_WEST", "Other"]
    assert zones["unit"].to_list() == ["GW", "GW"] and zones["mw"].round(1).to_list() == [1852.0, 1662.0]
    assert set(df["extraction_method"]) == {"pptx_chart"} and set(df["verified"]) == {False}
    assert df["chart_title"][0] == "ERCOT Approvals - Past 12 Months"


def test_datasets_are_declared():
    assert {d.name: d.mode for d in DATASETS} == {
        "large_load_headlines": "by_file", "large_load_status": "by_file", "document_pages": "by_file"}
