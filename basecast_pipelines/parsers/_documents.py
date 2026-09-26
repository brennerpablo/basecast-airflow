"""Shared helpers for document sources (ERCOT decks, PUCT filings): page text from PDF and PPTX, personal
data scrubbing, native PPTX tables and charts, and the ``document_pages`` table that both sources write.

Text comes from the text layer only (pdfplumber for PDF, python-pptx for PPTX): no OCR, so scanned pages
have no text. Extraction stops after ``MAX_PAGES`` pages per document to bound runtime; rows say so in
``truncated``. Extracted page text is cached per raw file (key + sha256) for the life of the process, so
several datasets reading the same documents parse each one once.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from collections import OrderedDict
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import pdfplumber
import polars as pl
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

log = logging.getLogger(__name__)

MAX_PAGES = 60
_CACHE_MAX_CHARS = 60_000_000  # ~120 MB of str at most; the full PUCT corpus is ~10M chars

# --- personal data -----------------------------------------------------------------------------------

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# US phone numbers written with dashes, dots or a parenthesized area code. Plain space-separated digit
# runs are left alone: chart axes in the text layer look like "250 500 1000".
_PHONE = re.compile(r"(?<![\w-])(?:\+?1[\s.-]?)?(?:\(\d{3}\)\s?|\d{3}[.-])\d{3}[.-]\d{4}(?![\w-])")
_BAR_NUMBER = re.compile(r"(?i)(state\s+bar\s+(?:no\.?|number)\s*:?\s*)\d{5,9}")


def scrub(text: str | None) -> str | None:
    """Drop emails, phone numbers and bar numbers from free text (they identify individuals)."""
    if not text:
        return text
    text = _EMAIL.sub("[email]", text)
    text = _PHONE.sub("[phone]", text)
    return _BAR_NUMBER.sub(r"\1[number]", text)


# --- text normalization --------------------------------------------------------------------------------

_DASHES = str.maketrans({"–": "-", "—": "-", "‑": "-", " ": " ", " ": " ", "\x0b": "\n"})


def clean_text(text: str | None) -> str:
    """Normalize dashes and odd spaces, drop NULs; keep line breaks."""
    if not text:
        return ""
    text = text.translate(_DASHES).replace("\x00", "")
    lines = [re.sub(r"[ \t\f\r]+", " ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def flatten(text: str) -> str:
    """One line with single spaces (regex matching across wrapped lines)."""
    return re.sub(r"\s+", " ", text).strip()


_BOILERPLATE_LINE = re.compile(r"^(?:ERCOT\s+)?PUBLIC(?:\s+\d+)?$|^\d{1,3}$|^Item\s+[\d.]+$|^ercot\W*$", re.IGNORECASE)


def page_title(text: str) -> str | None:
    """First meaningful line of a page (slide title), skipping footers such as "PUBLIC 5"."""
    for line in text.split("\n"):
        line = line.strip()
        if line and not _BOILERPLATE_LINE.match(line) and re.search(r"[A-Za-z]{3}", line):
            return line[:200]
    return None


# --- documents -------------------------------------------------------------------------------------------


@dataclass
class Page:
    number: int  # 1-based page or slide number
    text: str  # cleaned, not yet scrubbed
    title: str | None = None


@dataclass
class ExtractedDoc:
    kind: str  # "pdf" | "pptx"
    page_count: int
    pages: list[Page] = field(default_factory=list)
    truncated: bool = False
    error: str | None = None

    @property
    def n_chars(self) -> int:
        return sum(len(p.text) for p in self.pages)


_CACHE: OrderedDict[tuple[str, str, str], ExtractedDoc] = OrderedDict()
_CACHE_CHARS = 0


def _cached(key: tuple[str, str, str], make) -> ExtractedDoc:
    global _CACHE_CHARS
    if key in _CACHE:
        _CACHE.move_to_end(key)
        return _CACHE[key]
    doc = make()
    _CACHE[key] = doc
    _CACHE_CHARS += doc.n_chars
    while _CACHE_CHARS > _CACHE_MAX_CHARS and len(_CACHE) > 1:
        _, old = _CACHE.popitem(last=False)
        _CACHE_CHARS -= old.n_chars
    return doc


def pdf_document(source: Path | bytes, *, max_pages: int = MAX_PAGES) -> ExtractedDoc:
    """Page text of a PDF (text layer only). Errors are recorded on the result, not raised: a damaged page
    or file must not stop a corpus run; callers that need the text check ``error``."""
    doc = ExtractedDoc(kind="pdf", page_count=0)
    try:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        with pdfplumber.open(stream) as pdf:
            doc.page_count = len(pdf.pages)
            doc.truncated = doc.page_count > max_pages
            for i, page in enumerate(pdf.pages[:max_pages], 1):
                try:
                    text = clean_text(page.extract_text())
                except Exception as exc:  # noqa: BLE001 - pdfminer raises many kinds on bad pages
                    log.warning("page %d: text extraction failed (%s)", i, exc)
                    doc.error = f"page {i}: {type(exc).__name__}: {exc}"[:300]
                    text = ""
                finally:
                    page.close()
                doc.pages.append(Page(i, text, page_title(text)))
    except Exception as exc:  # noqa: BLE001
        doc.error = f"{type(exc).__name__}: {exc}"[:300]
    return doc


def _shape_texts(shapes) -> Iterator[str]:
    """Text of every shape in reading order, entering groups; tables row by row ("a | b | c")."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes)
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            yield shape.text_frame.text
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(clean_text(c.text).replace("\n", " ") for c in row.cells)


def _slide_title(slide) -> str | None:
    title = slide.shapes.title
    if title is not None and title.has_text_frame and title.text_frame.text.strip():
        return flatten(clean_text(title.text_frame.text))[:200]
    return None


def pptx_document(data: bytes, *, max_pages: int = MAX_PAGES) -> ExtractedDoc:
    """Slide text of a PPTX (shapes, groups and tables; speaker notes are left out)."""
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    doc = ExtractedDoc(kind="pptx", page_count=len(slides), truncated=len(slides) > max_pages)
    for i, slide in enumerate(slides[:max_pages], 1):
        text = clean_text("\n".join(_shape_texts(slide.shapes)))
        doc.pages.append(Page(i, text, _slide_title(slide) or page_title(text)))
    return doc


def extract(kind: str, data_or_path: Path | bytes, *, cache_key: tuple[str, str, str] | None = None) -> ExtractedDoc:
    """Extract one document; ``cache_key`` = (raw key, sha256, zip member or "")."""

    def make() -> ExtractedDoc:
        if kind == "pdf":
            return pdf_document(data_or_path)
        data = data_or_path if isinstance(data_or_path, bytes) else Path(data_or_path).read_bytes()
        return pptx_document(data)

    return _cached(cache_key, make) if cache_key else make()


def doc_kind(name: str) -> str | None:
    lowered = name.lower()
    if lowered.endswith(".pdf"):
        return "pdf"
    if lowered.endswith(".pptx"):
        return "pptx"
    return None


def zip_members(data: bytes) -> list[zipfile.ZipInfo]:
    return [i for i in zipfile.ZipFile(io.BytesIO(data)).infolist() if not i.is_dir()]


# --- PPTX slides with native tables and charts --------------------------------------------------------


def iter_shapes(shapes) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def pptx_slides(data: bytes) -> list[Any]:
    return list(Presentation(io.BytesIO(data)).slides)


def table_grid(shape) -> list[list[str]]:
    """A native PPTX table as rows of cleaned cell text."""
    return [[flatten(clean_text(c.text)) for c in row.cells] for row in shape.table.rows]


def slide_title(slide) -> str | None:
    title = _slide_title(slide)
    if title:
        return title
    texts = [flatten(clean_text(t)) for t in _shape_texts(slide.shapes)]
    return next((t[:200] for t in texts if re.search(r"[A-Za-z]{3}", t)), None)


# --- the shared document_pages table ---------------------------------------------------------------

PAGES_SCHEMA: dict[str, pl.DataType] = {
    "source_id": pl.String,
    "doc_type": pl.String,
    "file_name": pl.String,
    "member": pl.String,
    "title": pl.String,
    "document_date": pl.Date,
    "docket": pl.String,
    "item": pl.Int32,
    "url": pl.String,
    "page": pl.Int32,
    "page_title": pl.String,
    "text": pl.String,
    "n_chars": pl.Int32,
    "page_count": pl.Int32,
    "truncated": pl.Boolean,
}

PAGES_DESCRIPTION = (
    "Page/slide text of ERCOT large-load decks and PUCT filings (text layer only, no OCR; first "
    f"{MAX_PAGES} pages per document, `truncated` marks longer ones). Emails and phone numbers are "
    "replaced by [email]/[phone]; filings by individuals are left out. Evidence text for the commercial "
    "module. One row per page with text."
)


@dataclass
class DocInfo:
    source_id: str
    doc_type: str | None
    file_name: str
    member: str | None
    title: str | None
    document_date: date | None
    url: str | None
    docket: str | None = None
    item: int | None = None


def page_rows(info: DocInfo, doc: ExtractedDoc) -> list[dict]:
    rows = []
    for page in doc.pages:
        text = scrub(page.text)
        if not text:
            continue
        rows.append({
            "source_id": info.source_id,
            "doc_type": info.doc_type,
            "file_name": info.file_name,
            "member": info.member,
            "title": info.title,
            "document_date": info.document_date,
            "docket": info.docket,
            "item": info.item,
            "url": info.url,
            "page": page.number,
            "page_title": scrub(page.title),
            "text": text,
            "n_chars": len(text),
            "page_count": doc.page_count,
            "truncated": doc.truncated,
        })
    return rows


def frame(rows: list[dict], schema: dict[str, pl.DataType]) -> pl.DataFrame | None:
    """Rows → DataFrame with the declared schema (so an all-null column keeps its type in Postgres)."""
    if not rows:
        return None
    return pl.DataFrame(rows, schema=schema, orient="row" if isinstance(rows[0], (list, tuple)) else None)
